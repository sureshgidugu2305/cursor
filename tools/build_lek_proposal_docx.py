"""Generate the LEK Agentforce enhancement proposal as a Word document.

Run: python3 tools/build_lek_proposal_docx.py [output_path]
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x0B, 0x2B, 0x45)
ACCENT = RGBColor(0x00, 0x6D, 0xCC)
GREY = RGBColor(0x5A, 0x63, 0x6A)
HEADER_FILL = "0B2B45"
BAND_FILL = "EEF3F8"

DOC_TITLE = "Agentforce Enhancement Proposal"
DOC_SUBTITLE = "From CRM Record Assistant to Governed AI Sales Coworker"
CLIENT = "LEAK (LEK)"
PREPARED_BY = "Salesforce Architecture Practice"
VERSION = "Version 1.0"


def set_cell_background(cell, hex_fill):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shading)


def style_base(document):
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before in (
        ("Heading 1", 18, NAVY, 18),
        ("Heading 2", 13.5, NAVY, 14),
        ("Heading 3", 11.5, ACCENT, 10),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True


def add_footer(document):
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.text = f"{CLIENT} | {DOC_TITLE} | Confidential"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.runs[0]
    run.font.size = Pt(8)
    run.font.color.rgb = GREY


def add_paragraph(document, text, *, italic=False, size=10.5, color=None, space_after=6):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    run = paragraph.add_run(text)
    run.italic = italic
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return paragraph


def add_bullets(document, items, style="List Bullet"):
    for item in items:
        paragraph = document.add_paragraph(style=style)
        paragraph.paragraph_format.space_after = Pt(3)
        if isinstance(item, tuple):
            lead, rest = item
            bold_run = paragraph.add_run(lead)
            bold_run.bold = True
            paragraph.add_run(rest)
        else:
            paragraph.add_run(item)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(cell, HEADER_FILL)

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(row):
            cell = cells[col_index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run(str(value))
            run.font.size = Pt(9.5)
            if col_index == 0:
                run.bold = True
            if row_index % 2 == 1:
                set_cell_background(cell, BAND_FILL)

    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)

    document.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def add_callout(document, title, body):
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    set_cell_background(cell, BAND_FILL)
    cell.text = ""

    heading = cell.paragraphs[0]
    heading_run = heading.add_run(title)
    heading_run.bold = True
    heading_run.font.size = Pt(10.5)
    heading_run.font.color.rgb = NAVY

    body_paragraph = cell.add_paragraph()
    body_run = body_paragraph.add_run(body)
    body_run.font.size = Pt(10)

    document.add_paragraph().paragraph_format.space_after = Pt(4)


def build_cover(document):
    for _ in range(4):
        document.add_paragraph()

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(DOC_TITLE)
    title_run.bold = True
    title_run.font.size = Pt(30)
    title_run.font.color.rgb = NAVY

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(DOC_SUBTITLE)
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.color.rgb = ACCENT

    client = document.add_paragraph()
    client.alignment = WD_ALIGN_PARAGRAPH.CENTER
    client_run = client.add_run(f"Prepared for {CLIENT}")
    client_run.font.size = Pt(12)
    client_run.bold = True

    for _ in range(6):
        document.add_paragraph()

    meta = add_table(
        document,
        ["Proposal detail", "Value"],
        [
            ["Client", CLIENT],
            ["Prepared by", PREPARED_BY],
            ["Document", f"{DOC_TITLE} — {VERSION}"],
            ["Current baseline", "Agentforce implemented mid-2025"],
            ["Comparison platform", "Agentforce 360 (current release family)"],
            ["Scope in place today", "Account, Contact, Opportunity creation; record identification and querying"],
            ["Proposed outcome", "Governed sales agent that researches, briefs, updates, and follows up"],
        ],
        widths=[2.2, 4.3],
    )
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_run = note.add_run(
        "Confidential. Feature availability and licensing must be validated in the LEAK production org before build."
    )
    note_run.italic = True
    note_run.font.size = Pt(8.5)
    note_run.font.color.rgb = GREY

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_document(output_path):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    style_base(document)
    add_footer(document)
    build_cover(document)

    # 1. Executive summary
    document.add_heading("1. Executive Summary", level=1)
    add_paragraph(
        document,
        "LEAK implemented Agentforce in mid-2025 and proved the core value case: an employee agent that identifies "
        "Salesforce records, answers questions about CRM data, and creates Account, Contact, and Opportunity records "
        "from conversation. That was the right first release for the platform capabilities available at the time.",
    )
    add_paragraph(
        document,
        "Agentforce has advanced substantially since then. Current releases deliver hybrid reasoning with deterministic "
        "control, prebuilt sales research, grounding in unstructured firm content, automated testing, and production "
        "observability. The practical effect is that an Agentforce agent is no longer limited to creating records — it "
        "can prepare partners for meetings, keep pipeline current, draft grounded outreach, and answer from LEAK "
        "methodology content, all under human approval where commercial risk exists.",
    )
    add_callout(
        document,
        "Proposal in one sentence",
        "Extend LEAK's existing Agentforce investment from a CRM record assistant into a governed AI sales coworker by "
        "completing the Account / Contact / Opportunity lifecycle, adding research and knowledge grounding, and "
        "introducing deterministic control with testing and observability.",
    )
    add_paragraph(document, "We propose a three-phase programme:")
    add_bullets(
        document,
        [
            ("Phase 1 — Trusted CRM coworker: ", "summarize, update, and task actions on the objects LEAK already creates, protected by Agent Script, confirmations, and automated tests."),
            ("Phase 2 — Research and knowledge: ", "account research briefings, an Agentforce Data Library of LEAK playbooks and templates, grounded email drafting, and Slack access."),
            ("Phase 3 — Growth: ", "prospecting signals, Sales Workspace visibility, multi-agent handoffs, and longer-running business development goals where justified."),
        ],
    )

    # 2. Current state
    document.add_heading("2. Current State — What LEAK Has Today", level=1)
    add_paragraph(
        document,
        "The existing implementation is a single employee-facing agent operating inside Salesforce with the mid-2025 "
        "topic-and-action model.",
    )
    add_table(
        document,
        ["Capability", "Status today", "Practical effect for users"],
        [
            ["Identify record by name", "In place", "Resolves a named company or person to the correct record"],
            ["Query records", "In place", "Answers questions about data already held in CRM"],
            ["Create Account", "In place", "New client or target organisation created from conversation"],
            ["Create Contact", "In place", "Stakeholder captured against the account"],
            ["Create Opportunity", "In place", "New pursuit opened without navigating the UI"],
            ["Update existing records", "Not in scope", "Stage, close date, and next step still maintained manually"],
            ["Account research and summaries", "Not in scope", "Partners prepare for meetings outside the agent"],
            ["Knowledge and playbook answers", "Not in scope", "Firm methodology content is not available to the agent"],
            ["Email drafting and activity logging", "Not in scope", "Outreach and follow-up remain fully manual"],
            ["Deterministic process control", "Not available in 2025 design", "Behaviour depends on prompt instructions alone"],
            ["Automated testing and observability", "Limited", "Difficult to measure quality or safely widen write access"],
        ],
        widths=[1.9, 1.5, 3.3],
    )
    add_callout(
        document,
        "Architectural assessment",
        "The agent is a data-entry and lookup assistant. Creation is solved; the remainder of the sales cycle — "
        "preparation, judgement, hygiene, and follow-up — is still entirely human. Most of the available value from the "
        "existing licence is therefore unrealised.",
    )

    document.add_page_break()

    # 3. What changed
    document.add_heading("3. What Has Changed in Agentforce Since Mid-2025", level=1)
    add_paragraph(
        document,
        "The following table summarises publicly announced Salesforce capability waves relevant to LEAK. Availability "
        "under LEAK's specific licence set must be confirmed during discovery.",
    )
    add_table(
        document,
        ["Release wave", "Timing", "Capabilities relevant to LEAK"],
        [
            ["Agentforce 3", "June 2025", "Command Center observability, Testing Center enhancements, 100+ additional prebuilt actions, MCP interoperability, web search in data libraries"],
            ["Agentforce 360", "October 2025", "New Agentforce Builder, Agent Script hybrid reasoning, Agentforce Voice, Intelligent Context for unstructured documents, Data 360 grounding"],
            ["Winter '26", "October 2025", "Agentforce Grid for bulk prompt and action experimentation, pipeline management enhancements, built-in observability"],
            ["Spring '26", "February 2026", "Sales Workspace, Account Management with automated research, always-on prospecting, conversational email"],
            ["Terminology change", "April 2026", "Topics are now called subagents, reflecting a clearer multi-skill agent design"],
            ["Summer '26", "June 2026", "Multi-agent orchestration with shared context, Slack-first agentic workflows, agentic data management"],
            ["Long-horizon agents", "2026 announcements", "Goal-oriented agents with memory and durable execution across days or weeks, plus Agent Optimizer for trace-driven tuning"],
        ],
        widths=[1.4, 1.1, 4.2],
    )
    add_paragraph(
        document,
        "The net shift is architectural. Mid-2025 agents were conversational interfaces over CRM. Current Agentforce is "
        "a controlled execution layer over Customer 360 and Data 360, with deterministic scripting, automated testing, "
        "and production monitoring.",
    )

    # 4. Gap analysis
    document.add_heading("4. Gap Analysis — Current Scope Against Current Platform", level=1)
    add_table(
        document,
        ["Domain", "LEAK today", "Available now", "Business benefit"],
        [
            ["CRM lifecycle", "Create only", "Get details, summarize, update, aggregate queries, field extraction", "Pipeline stays current without leaving the agent"],
            ["Account research", "Manual preparation", "Sales Research subagent: activities, notes, emails, account plan, governed web search", "Meeting preparation reduced from hours to minutes"],
            ["Firm knowledge", "CRM fields only", "Agentforce Data Libraries, Intelligent Context for documents and decks", "Consistent answers from LEAK playbooks and templates"],
            ["Communication", "Not enabled", "Grounded email drafting and revision, in-record recommended actions", "Faster, on-brand client outreach"],
            ["Activity capture", "Not enabled", "Activity timelines and task or event creation via Flow", "Follow-ups are actually recorded"],
            ["Process control", "Prompt instructions", "Agent Script with conditions, variables, and action chaining", "Duplicate checks and approvals enforced, not suggested"],
            ["Quality assurance", "Manual testing", "Testing Center, session traces, Command Center observability", "Write access can be widened safely"],
            ["Channels", "Salesforce only", "Slack, mobile, Sales Workspace, optional voice", "Adoption where partners already work"],
            ["Scope of a request", "Single turn", "Multi-step plans, bulk operations, long-horizon goals", "Larger outcomes per interaction"],
            ["Integration", "Salesforce only", "MCP, MuleSoft, external services", "Enrichment and downstream systems in the loop"],
        ],
        widths=[1.1, 1.3, 2.4, 1.9],
    )

    document.add_page_break()

    # 5. Proposed solution
    document.add_heading("5. Proposed Solution", level=1)
    add_paragraph(
        document,
        "We recommend retaining a single employee-facing agent as the user experience and structuring capability behind "
        "it as purpose-built subagents. This preserves a simple front door for partners while allowing each job to be "
        "instructed, tested, and governed independently.",
    )

    document.add_heading("5.1 Proposed subagent design", level=2)
    add_table(
        document,
        ["Subagent", "Purpose", "Representative actions"],
        [
            ["Record management", "Create and maintain core CRM records", "Identify, query, create (existing), update, extract fields and values"],
            ["Sales research", "Brief partners on an account or pursuit", "Record details, activity timeline, related notes and emails, account plan, governed web search, summarize"],
            ["Knowledge", "Answer questions from LEAK content", "Answer questions with knowledge, with citations to source material"],
            ["Outreach", "Support client communication and follow-up", "Draft or revise email, create task or event"],
            ["Pipeline insight", "Report on pursuit health", "Aggregate queries, at-risk and stale pipeline summaries"],
        ],
        widths=[1.3, 2.0, 3.4],
    )

    document.add_heading("5.2 Design principles", level=2)
    add_bullets(
        document,
        [
            ("Least privilege. ", "The agent operates with the same sharing, field-level security, and validation rules as the human user it serves."),
            ("Read freely, write with confirmation. ", "Explicit user approval is required for new opportunities, stage changes to closed, amounts above a threshold, and ownership changes."),
            ("Duplicate prevention before creation. ", "Identification and matching run deterministically in Agent Script before any record is created."),
            ("Ground every generative response. ", "Summaries, briefings, and email drafts must draw on CRM records and approved library content; no unsupported commercial claims."),
            ("One job per subagent. ", "Research, record maintenance, and knowledge are separated so each can be instructed and tested precisely."),
            ("Treat retrieved text as data. ", "Notes, emails, and web content are inputs to reasoning, never instructions to the agent."),
        ],
    )

    document.add_heading("5.3 Guardrails proposed for write actions", level=2)
    add_table(
        document,
        ["Control", "Applies to", "Mechanism"],
        [
            ["User confirmation", "Create opportunity, stage change, amount change, owner change", "Standard action confirmation plus Agent Script conditions"],
            ["Duplicate check", "All create actions", "Identify and query chain executed before data manipulation"],
            ["Required fields", "Opportunity and account creation", "Flow-enforced validation rather than prompt guidance"],
            ["Prohibited operations", "Delete and merge", "Excluded from the agent; retained in data stewardship process"],
            ["Auditability", "All actions", "Einstein Trust Layer and Command Center session traces"],
        ],
        widths=[1.3, 2.4, 3.0],
    )

    document.add_page_break()

    # 6. Advantages
    document.add_heading("6. Advantages to LEAK", level=1)

    document.add_heading("6.1 For partners and business development", level=2)
    add_table(
        document,
        ["Advantage", "Why it matters"],
        [
            ["Faster meeting preparation", "A single request produces an account briefing drawn from records, activity, and notes"],
            ["Reduced CRM administration", "Stage, next step, and follow-up tasks maintained conversationally"],
            ["Better client outreach", "Email drafts grounded in the specific opportunity and stakeholder context"],
            ["Access to firm intellectual property", "Playbooks, scope language, and templates answerable in the flow of work"],
            ["Available where work happens", "Salesforce, Slack, and mobile rather than a single interface"],
        ],
        widths=[2.2, 4.4],
    )

    document.add_heading("6.2 For CRM and data quality", level=2)
    add_table(
        document,
        ["Advantage", "Why it matters"],
        [
            ["Fewer duplicate records", "Identification and matching enforced deterministically before creation"],
            ["More complete opportunities", "Update and task actions close the gap left by create-only scope"],
            ["Traceable agent activity", "Every action visible through observability and the trust layer"],
            ["Controlled commercial writes", "Confirmation required on the fields that affect reported pipeline"],
        ],
        widths=[2.2, 4.4],
    )

    document.add_heading("6.3 For IT and platform ownership", level=2)
    add_table(
        document,
        ["Advantage", "Why it matters"],
        [
            ["Predictable behaviour", "Agent Script provides deterministic logic without abandoning generative flexibility"],
            ["Less custom build", "Prebuilt research and account management capability replaces bespoke development"],
            ["Safe change management", "Testing Center supports regression testing before production release"],
            ["Operational visibility", "Command Center exposes adoption, failure patterns, latency, and cost"],
            ["Controlled extensibility", "MCP and MuleSoft used only where native Salesforce actions are insufficient"],
        ],
        widths=[2.2, 4.4],
    )

    document.add_heading("6.4 For LEAK leadership", level=2)
    add_table(
        document,
        ["Advantage", "Why it matters"],
        [
            ["Higher return on existing investment", "The same platform licence covers research, hygiene, and outreach, not creation alone"],
            ["Measurable outcomes", "Preparation time, duplicate rate, pipeline completeness, and action success are all reportable"],
            ["Scalable growth path", "Prospecting and longer-running goals can be added without re-platforming"],
            ["Risk-managed autonomy", "Agent independence expands only where tests and confirmations demonstrate reliability"],
        ],
        widths=[2.2, 4.4],
    )

    # 7. Roadmap
    document.add_heading("7. Proposed Delivery Roadmap", level=1)
    add_paragraph(
        document,
        "Phasing is sequenced by risk. Control and quality capability is introduced alongside the first functional "
        "expansion so that later phases can be delivered with confidence.",
    )
    add_table(
        document,
        ["Phase", "Scope", "Exit criteria"],
        [
            ["Phase 0 — Discovery", "Inventory existing topics, actions, flows, and permissions; confirm licensing; establish baseline metrics", "Agreed scope, confirmed feature availability, measured starting position"],
            ["Phase 1 — Trusted CRM coworker", "Record details, summarize, confirmed update, activity timeline, task creation, Agent Script duplicate prevention, Testing Center, observability, record-page actions", "Partners can brief and update an opportunity in the agent; no unconfirmed closed-won changes; duplicate creation rate reduced"],
            ["Phase 2 — Research and knowledge", "Sales research subagent, Agentforce Data Library of playbooks and templates, grounded email drafting without automatic send, Slack channel", "Preparation time materially reduced; knowledge answers cited; email drafts accepted with light editing"],
            ["Phase 3 — Growth", "Prospecting signals, Sales Workspace, multi-agent handoff where a second domain exists, evaluation of long-horizon goals and Agent Optimizer", "Qualified pipeline supported proactively; leadership visibility of agent contribution"],
        ],
        widths=[1.4, 2.7, 2.5],
    )
    add_callout(
        document,
        "Sequencing recommendation",
        "Approve Phase 1 now and run Data Library scoping in parallel. Defer prospecting, voice, and multi-agent "
        "orchestration until Phase 1 quality metrics are green.",
    )

    # 8. Metrics
    document.add_heading("8. Success Measures", level=1)
    add_table(
        document,
        ["Measure", "Purpose"],
        [
            ["Weekly active agent users by role", "Adoption across partners and coordinators"],
            ["Time from new pursuit to complete opportunity record", "Process efficiency"],
            ["Duplicate account, contact, and opportunity rate", "Data quality"],
            ["Proportion of opportunities with a next step and future task", "Pipeline hygiene"],
            ["Knowledge citation rate and negative feedback rate", "Grounding quality"],
            ["Action success rate and escalation rate", "Reliability"],
            ["Human edit effort on drafted email", "Generative output quality"],
        ],
        widths=[2.8, 3.8],
    )

    # 9. Risks
    document.add_heading("9. Prerequisites, Risks, and Mitigations", level=1)
    add_table(
        document,
        ["Item", "Risk if unaddressed", "Mitigation"],
        [
            ["CRM data quality", "The agent amplifies existing duplicates and incomplete records", "Data remediation and matching rules during Phase 0 and 1"],
            ["Licence coverage", "Design depends on capability that cannot be enabled", "Confirm entitlements in discovery before build commitment"],
            ["Sharing model", "Sensitive engagement data over-exposed through querying", "Least-privilege agent user and review of sharing rules"],
            ["Client confidentiality", "Restricted content indexed into a knowledge library", "Content classification and retention rules before library load"],
            ["Prompt injection through retrieved text", "Untrusted notes or web content influence agent behaviour", "Research subagent treats retrieved content strictly as data"],
            ["Change control", "Instruction drift degrades production behaviour", "Sandbox versioning and regression tests before release"],
        ],
        widths=[1.4, 2.6, 2.6],
    )

    # 10. Recommendation
    document.add_heading("10. Recommendation and Next Steps", level=1)
    add_paragraph(
        document,
        "LEAK's mid-2025 implementation proved that Agentforce can reliably create and retrieve CRM records. The current "
        "platform makes it realistic to operate a governed sales coworker that also researches accounts, maintains "
        "pipeline, drafts client communication, and answers from firm knowledge, while humans retain approval over "
        "commercial writes.",
    )
    add_paragraph(document, "We recommend the following immediate steps:")
    add_bullets(
        document,
        [
            ("Approve Phase 1. ", "Complete the record lifecycle on Account, Contact, and Opportunity with Agent Script control and automated testing."),
            ("Commission discovery. ", "Confirm licensing, review existing topics and actions, and establish baseline metrics."),
            ("Begin content scoping. ", "Identify the playbooks, templates, and methodology documents suitable for a governed knowledge library."),
            ("Defer expansion. ", "Hold prospecting, voice, and multi-agent orchestration until Phase 1 quality measures are demonstrated."),
        ],
        style="List Number",
    )
    add_callout(
        document,
        "Decision requested",
        "Approval to proceed with Phase 0 discovery and Phase 1 delivery, with Data Library scoping running in parallel "
        "in preparation for Phase 2.",
    )

    add_paragraph(
        document,
        "Prepared as a Salesforce architecture recommendation. Capability availability, release status, and licensing "
        "are subject to validation in the LEAK org before implementation planning and estimation.",
        italic=True,
        size=9,
        color=GREY,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("docs/LEAK-Agentforce-Enhancement-Proposal.docx")
    saved = build_document(target)
    print(f"Saved {saved} ({saved.stat().st_size / 1024:.1f} KB)")
